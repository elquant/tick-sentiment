import settings
import pandas as pd

from datetime import datetime
from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference

tweets = pd.read_csv(settings.TWEETS_FILE_CSV)
tweets["user_created"] = tweets["user_created"].apply(lambda x: datetime.strptime(x, '%Y-%m-%dT%H:%M:%S'))
positive_polarity = tweets[tweets["polarity"] >= 0]

user_followers_by_year_month = pd.pivot_table(positive_polarity,
                                              index=positive_polarity['user_created'].apply(lambda x: x.sheet_),
                                              columns=positive_polarity["user_created"].apply(lambda x: x.year),
                                              values='user_followers',
                                              aggfunc='sum')

print("Pivot table for the total sum of followers grouped by year and month:")
print(user_followers_by_year_month.head())

file_path = settings.TWEETS_FILE_EXCEL
sheet_name = 'Followers by year and month'
user_followers_by_year_month.to_excel(file_path, sheet_name=sheet_name, startrow=3)

wb = load_workbook(file_path)
sheet1 = wb[sheet_name]

sheet1['A1'] = 'Followers by year and month'
sheet1['A2'] = 'elquant.com'
sheet1['A4'] = 'Month/Year'
sheet1['A1'].style = 'Title'
sheet1['A2'].style = 'Headline 2'

bar_chart = BarChart()
data = Reference(sheet1, min_col=2, max_col=15, min_row=4, max_row=16)
categories = Reference(sheet1, min_col=1, max_col=1, min_row=5, max_row=16)
bar_chart.add_data(data, titles_from_data=True)
bar_chart.set_categories(categories)
sheet1.add_chart(bar_chart, anchor="B19")
bar_chart.title = 'Aggregated number of followers by year and month'
bar_chart.style = 3

wb.save(filename=file_path)
